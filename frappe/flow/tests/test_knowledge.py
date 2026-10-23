# Copyright (c) 2026, Frappe Technologies and contributors
# License: MIT. See LICENSE

import json
from types import SimpleNamespace
from unittest.mock import patch

import frappe
from frappe.flow import ToolContext
from frappe.flow.knowledge import Knowledge, attachment_store, store
from frappe.flow.knowledge.chunker import chunk_text
from frappe.flow.knowledge.embedder import embed_texts, probe_dimension
from frappe.flow.knowledge.extract import (
	_extract_by_extension,
	_extract_docx,
	_extract_html,
	_extract_pdf,
	_extract_pdf_page,
	_extract_xlsx,
	_render_rich_text,
	_validate_public_url,
	extract,
)
from frappe.flow.knowledge.ingest import ingest_source, reconcile_source, sync_due_sources
from frappe.flow.knowledge.retriever import retrieve, retrieve_attachments
from frappe.tests import IntegrationTestCase

DIM = 4
CTX = ToolContext(session_id="test", agent="Test Agent", user="Administrator")


def _rows():
	return [
		{
			"id": 1,
			"kb": "Helpdesk",
			"source": "SRC-a",
			"content": "wifi not connecting on macbook",
			"vector": [1.0, 0.0, 0.0, 0.0],
		},
		{
			"id": 2,
			"kb": "Helpdesk",
			"source": "SRC-a",
			"content": "printer not detected on windows",
			"vector": [0.0, 1.0, 0.0, 0.0],
		},
		{
			"id": 3,
			"kb": "Onboarding",
			"source": "SRC-b",
			"content": "how to set up your laptop",
			"vector": [0.9, 0.1, 0.0, 0.0],
		},
	]


class TestKnowledgeStore(IntegrationTestCase):
	def setUp(self):
		store.drop_table()

	def tearDown(self):
		store.drop_table()

	def _seed(self):
		store.ensure_table_for_dimension(DIM)
		store.add(_rows())

	def test_the_table_is_created_once_for_one_dimension(self):
		self.assertEqual(store.search([1.0, 0.0, 0.0, 0.0]), [])
		with self.assertRaisesRegex(frappe.ValidationError, "not initialized"):
			store.add(_rows())
		with self.assertRaises(ValueError):
			store.ensure_table_for_dimension(0)

		store.ensure_table_for_dimension(DIM)
		store.ensure_table_for_dimension(DIM)
		self.assertTrue(store.table_exists())
		self.assertEqual(store.table_dimension(), DIM)
		self.assertEqual(store.search([1.0, 0.0, 0.0, 0.0]), [])
		with self.assertRaisesRegex(frappe.ValidationError, "Rebuild"):
			store.ensure_table_for_dimension(DIM + 1)

	def test_search_ranks_scopes_and_matches_keywords(self):
		self._seed()
		hits = store.search([1.0, 0.0, 0.0, 0.0], limit=2)
		self.assertEqual([h["id"] for h in hits], [1, 3])
		self.assertGreater(hits[0]["score"], hits[1]["score"])

		hits = store.search([1.0, 0.0, 0.0, 0.0], kbs=["Onboarding"], limit=5)
		self.assertEqual([h["kb"] for h in hits], ["Onboarding"])
		# a quote in a filter value cannot break out of the predicate
		self.assertEqual(store.search([1.0, 0.0, 0.0, 0.0], kbs=["x'; DROP TABLE chunks; --"], limit=5), [])
		self.assertTrue(store.table_exists())

		hits = store.search([0.0, 0.0, 1.0, 0.0], text="printer", limit=3)
		self.assertIn(2, [h["id"] for h in hits])
		# rows added after the full-text index was built are searched too
		store.add(
			[
				{
					"id": 4,
					"kb": "Helpdesk",
					"source": "SRC-c",
					"content": "vpn conflict after update",
					"vector": [0.0, 0.0, 1.0, 0.0],
				}
			]
		)
		hits = store.search([0.0, 0.0, 1.0, 0.0], text="vpn", limit=3)
		self.assertEqual(hits[0]["id"], 4)

	def test_delete_by_source_or_ids(self):
		with self.assertRaises(ValueError):
			store.delete()
		for criteria, remaining in (
			({"source": "SRC-a"}, [3]),
			({"ids": [1, 3]}, [2]),
			({"ids": []}, [1, 2, 3]),
		):
			with self.subTest(criteria=criteria):
				store.drop_table()
				self._seed()
				store.delete(**criteria)
				hits = store.search([1.0, 0.0, 0.0, 0.0], limit=5)
				self.assertEqual(sorted(h["id"] for h in hits), remaining)


def _embedding_response(vectors, shuffle=False):
	data = [{"embedding": v, "index": i} for i, v in enumerate(vectors)]
	if shuffle:
		data = list(reversed(data))
	return SimpleNamespace(data=data)


def _make_model(title="Embed Model", model_id="openai/text-embedding-3-small"):
	if frappe.db.exists("Flow Model", title):
		return frappe.get_doc("Flow Model", title)
	return frappe.get_doc(
		{"doctype": "Flow Model", "title": title, "model_id": model_id, "api_key": "sk-test", "enabled": 1}
	).insert()


def _set_settings(**values):
	for fieldname, value in values.items():
		frappe.db.set_single_value("Flow Knowledge Settings", fieldname, value)
	frappe.clear_document_cache("Flow Knowledge Settings", "Flow Knowledge Settings")


class TestEmbedder(IntegrationTestCase):
	def setUp(self):
		self.model = _make_model()
		_set_settings(embedding_model=self.model.name)

	def tearDown(self):
		frappe.db.rollback()

	def test_texts_are_embedded_in_order_with_the_models_config(self):
		self.model.db_set("extra_params", '{"provider": {"zdr": true}}')
		response = _embedding_response([[1.0, 2.0], [3.0, 4.0]], shuffle=True)
		with patch("litellm.embedding", return_value=response) as mocked:
			vectors = embed_texts(["a", "b"])
		self.assertEqual(vectors, [[1.0, 2.0], [3.0, 4.0]])
		kwargs = mocked.call_args.kwargs
		self.assertEqual(kwargs["model"], "openai/text-embedding-3-small")
		self.assertEqual(kwargs["api_key"], "sk-test")
		self.assertEqual(kwargs["input"], ["a", "b"])
		self.assertEqual(kwargs["provider"], {"zdr": True})

		def fake_embedding(input, **kwargs):
			return _embedding_response([[0.0]] * len(input))

		with patch("litellm.embedding", side_effect=fake_embedding) as mocked:
			self.assertEqual(len(embed_texts([f"t{i}" for i in range(100)])), 100)
		self.assertEqual([len(call.kwargs["input"]) for call in mocked.call_args_list], [96, 4])

		with patch("litellm.embedding") as mocked:
			self.assertEqual(embed_texts([]), [])
		mocked.assert_not_called()

		with patch("litellm.embedding", return_value=_embedding_response([[0.0] * 1536])):
			self.assertEqual(probe_dimension(self.model.name), 1536)

	def test_embedding_failures_surface_as_validation_errors(self):
		with (
			self.subTest("count mismatch"),
			patch("litellm.embedding", return_value=_embedding_response([[1.0]])),
			self.assertRaisesRegex(frappe.ValidationError, "2 inputs"),
		):
			embed_texts(["a", "b"])
		with (
			self.subTest("provider error"),
			patch("litellm.embedding", side_effect=RuntimeError("invalid api key")),
			self.assertRaisesRegex(frappe.ValidationError, "invalid api key"),
		):
			embed_texts(["a"])

		self.model.enabled = 0
		self.model.save()
		with self.subTest("disabled model"), self.assertRaisesRegex(frappe.ValidationError, "disabled"):
			embed_texts(["a"])

		_set_settings(embedding_model="")
		with self.subTest("no model"), self.assertRaisesRegex(frappe.ValidationError, "Knowledge Settings"):
			embed_texts(["a"])


class TestKnowledgeSettings(IntegrationTestCase):
	def setUp(self):
		self.model = _make_model()
		self.alt_model = _make_model(title="Embed Model Alt", model_id="openai/text-embedding-3-large")
		_set_settings(embedding_model="", embedding_dimension=0)

	def tearDown(self):
		frappe.db.rollback()

	def _save_settings(self, **values):
		settings = frappe.get_doc("Flow Knowledge Settings")
		settings.update(values)
		settings.save()
		return settings

	def _make_chunk(self):
		kb = frappe.get_doc({"doctype": "Flow Knowledge Base", "title": "Settings KB"}).insert()
		source = frappe.get_doc(
			{
				"doctype": "Flow Knowledge Source",
				"knowledge_base": kb.name,
				"source_type": "Text",
				"title": "Settings Source",
				"content": "hello",
			}
		).insert()
		return frappe.get_doc(
			{
				"doctype": "Flow Knowledge Chunk",
				"knowledge_base": kb.name,
				"source": source.name,
				"chunk_index": 0,
				"content": "hello",
			}
		).insert()

	def test_saving_probes_the_dimension_when_the_model_changes(self):
		with patch("frappe.flow.knowledge.embedder.probe_dimension", return_value=1536) as probe:
			settings = self._save_settings(embedding_model=self.model.name)
		self.assertEqual(settings.embedding_dimension, 1536)
		probe.assert_called_once_with(self.model.name)

		with patch("frappe.flow.knowledge.embedder.probe_dimension") as probe:
			self._save_settings(chunk_size=500)
		probe.assert_not_called()

		with patch("frappe.flow.knowledge.embedder.probe_dimension", return_value=3072):
			settings = self._save_settings(embedding_model=self.alt_model.name)
		self.assertEqual(settings.embedding_dimension, 3072)

		self.assertEqual(self._save_settings(embedding_model="").embedding_dimension, 0)

	def test_the_model_cannot_change_while_chunks_exist(self):
		with patch("frappe.flow.knowledge.embedder.probe_dimension", return_value=1536):
			self._save_settings(embedding_model=self.model.name)
		self._make_chunk()
		for model in (self.alt_model.name, ""):
			with (
				self.subTest(embedding_model=model),
				self.assertRaisesRegex(frappe.ValidationError, "Delete all knowledge sources"),
			):
				self._save_settings(embedding_model=model)

		# a rename keeps the model and its dimension
		with patch("frappe.flow.knowledge.embedder.probe_dimension") as probe:
			# Frappe trims the name
			frappe.rename_doc("Flow Model", self.model.name, " Embed Model Renamed ", rebuild_search=False)
		probe.assert_not_called()
		settings = frappe.get_doc("Flow Knowledge Settings")
		self.assertEqual(
			(settings.embedding_model, settings.embedding_dimension), ("Embed Model Renamed", 1536)
		)

		# merging it into another model is still a change
		with self.assertRaisesRegex(frappe.ValidationError, "Delete all knowledge sources"):
			frappe.rename_doc("Flow Model", "Embed Model Renamed", self.alt_model.name, merge=True)

		# a rebuild's flag lifts the guard
		settings = frappe.get_doc("Flow Knowledge Settings")
		settings.embedding_model = self.alt_model.name
		settings.flags.allow_embedding_model_change = True
		with patch("frappe.flow.knowledge.embedder.probe_dimension", return_value=3072):
			settings.save()
		self.assertEqual(settings.embedding_dimension, 3072)

	def test_chunk_size_and_overlap_are_validated(self):
		for values, error in (
			({"chunk_size": 100, "chunk_overlap": 100}, "Chunk Overlap"),
			({"chunk_size": 0}, "Chunk Size"),
		):
			with self.subTest(values=values), self.assertRaisesRegex(frappe.ValidationError, error):
				self._save_settings(**values)


class TestChunker(IntegrationTestCase):
	def test_text_is_split_into_overlapping_chunks_within_size(self):
		self.assertEqual(chunk_text("   \n  ", chunk_size=100, overlap=10), [])
		self.assertEqual(chunk_text("hello world", chunk_size=100, overlap=10), ["hello world"])
		# no whitespace to break on: a hard split
		self.assertEqual([len(c) for c in chunk_text("a" * 250, chunk_size=100, overlap=0)], [100, 100, 50])

		words = [f"w{i}" for i in range(200)]
		chunks = chunk_text(" ".join(words), chunk_size=100, overlap=20)
		self.assertGreater(len(chunks), 1)
		self.assertTrue(all(len(c) <= 100 for c in chunks))
		joined = " ".join(chunks)
		self.assertTrue(all(w in joined for w in words))

		chunks = chunk_text(" ".join(f"token{i}" for i in range(100)), chunk_size=80, overlap=30)
		self.assertIn(chunks[0].split()[-1], chunks[1])

		for size, overlap in ((0, 0), (10, 10)):
			with self.subTest(size=size, overlap=overlap), self.assertRaises(ValueError):
				chunk_text("x", chunk_size=size, overlap=overlap)


class TestExtract(IntegrationTestCase):
	def tearDown(self):
		frappe.db.rollback()

	def test_text_and_file_sources_extract_one_document(self):
		docs = extract(frappe._dict(source_type="Text", content="  some knowledge  "))
		self.assertEqual(len(docs), 1)
		self.assertEqual(docs[0].text, "some knowledge")
		self.assertIsNone(docs[0].reference_name)
		self.assertEqual(extract(frappe._dict(source_type="Text", content="   ")), [])

		file_doc = frappe.get_doc(
			{"doctype": "File", "file_name": "note.txt", "content": "file body", "is_private": 1}
		).insert()
		docs = extract(frappe._dict(source_type="File", file=file_doc.file_url))
		self.assertEqual(docs[0].text, "file body")

		file_doc = frappe.get_doc(
			{"doctype": "File", "file_name": "data.bin", "content": "x", "is_private": 1}
		).insert()
		with self.assertRaisesRegex(frappe.ValidationError, "Unsupported file format"):
			extract(frappe._dict(source_type="File", file=file_doc.file_url))

	def test_a_doctype_source_extracts_each_matching_record_as_plain_text(self):
		marker = frappe.generate_hash(length=10)
		frappe.get_doc(
			{
				"doctype": "ToDo",
				"description": f"<p>{marker} Cannot <b>login</b></p><script>x=1</script>",
				"status": "Open",
			}
		).insert()
		frappe.get_doc(
			{"doctype": "ToDo", "description": f"{marker} second task", "status": "Cancelled"}
		).insert()
		source = frappe._dict(
			source_type="DocType",
			reference_doctype="ToDo",
			content_fields="description",
			filters=json.dumps({"status": "Open", "description": ["like", f"%{marker}%"]}),
		)
		docs = extract(source)
		self.assertEqual(len(docs), 1)
		self.assertEqual(docs[0].reference_doctype, "ToDo")
		self.assertTrue(docs[0].reference_name)
		self.assertIn("Cannot login", docs[0].text)
		self.assertNotIn("<", docs[0].text)
		self.assertNotIn("x=1", docs[0].text)

	def test_a_doctype_source_refuses_fields_and_filters_it_cannot_use(self):
		for doctype, content_fields, filters, error in (
			("ToDo", "description, not_a_field", None, "Unknown fields"),
			("User", "full_name, roles", None, "Child table fields"),
			("ToDo", "description", "not json", "valid JSON"),
		):
			source = frappe._dict(
				source_type="DocType",
				reference_doctype=doctype,
				content_fields=content_fields,
				filters=filters,
			)
			with self.subTest(error), self.assertRaisesRegex(frappe.ValidationError, error):
				extract(source)

	def test_rich_text_fields_render_as_plain_text(self):
		markdown = frappe._dict(fieldtype="Markdown Editor", options=None)
		out = _render_rich_text(markdown, "# Title\n\nSome **bold** [link](https://x.io) text")
		self.assertIn("Title", out)
		self.assertIn("bold", out)
		self.assertNotIn("**", out)
		self.assertNotIn("https://x.io", out)

		markdown_code = frappe._dict(fieldtype="Code", options="Markdown")
		self.assertEqual(_render_rich_text(markdown_code, "## Heading\n\ncontent"), "Heading content")

		code = "# a comment\nx = 1"
		self.assertEqual(_render_rich_text(frappe._dict(fieldtype="Code", options="Python"), code), code)

	def test_office_and_html_files_extract_to_text(self):
		import io

		import openpyxl
		from docx import Document

		workbook = openpyxl.Workbook()
		sheet = workbook.active
		sheet.title = "People"
		sheet.append(["Name", "Age"])
		sheet.append(["Alice", 30])
		buffer = io.BytesIO()
		workbook.save(buffer)
		text = _extract_xlsx(buffer.getvalue())
		self.assertIn("People", text)
		self.assertIn("Alice", text)

		document = Document()
		document.add_paragraph("First paragraph.")
		document.add_paragraph("")
		document.add_paragraph("Second paragraph.")
		buffer = io.BytesIO()
		document.save(buffer)
		self.assertEqual(_extract_docx(buffer.getvalue()), "First paragraph.\n\nSecond paragraph.")

		html = "<html><body><script>x=1</script><p>Hi</p> <p>there</p></body></html>"
		self.assertEqual(_extract_html(html), "Hi there")

	def test_pdfs_and_images_are_read_with_one_ocr_pass_per_page(self):
		import io

		from PIL import Image, ImageDraw
		from pypdf import PdfWriter

		writer = PdfWriter()
		writer.add_blank_page(width=72, height=72)
		buffer = io.BytesIO()
		writer.write(buffer)
		self.assertEqual(_extract_pdf(buffer.getvalue()), "")

		page = SimpleNamespace(
			images=[{"x0": 0, "top": 0, "x1": 100, "bottom": 100}] * 16,
			bbox=(0, 0, 612, 792),
			find_tables=lambda: [],
			extract_text=lambda: "Searchable text",
			crop=lambda bbox: object(),
		)
		with (
			patch("frappe.flow.knowledge.extract._render_page_png", return_value=b"page"),
			patch("frappe.flow.knowledge.extract._ocr_image", return_value="Text in images") as ocr,
		):
			text = _extract_pdf_page(page)
		self.assertIn("Searchable text", text)
		self.assertIn("Text in images", text)
		self.assertEqual(ocr.call_count, 1)

		def scan(size, at, words, **save):
			image = Image.new("RGB", size, "white")
			ImageDraw.Draw(image).text(at, words, fill="black")
			buffer = io.BytesIO()
			image.save(buffer, **save)
			return buffer.getvalue()

		png = scan((400, 120), (20, 45), "Hello World", format="PNG")
		self.assertIn("Hello", _extract_by_extension(png, "png"))
		pdf = scan((600, 200), (30, 80), "Scanned Document", format="PDF", resolution=100)
		self.assertIn("Scanned", _extract_pdf(pdf))

	def test_only_public_http_urls_are_fetched(self):
		for url, error in (
			("ftp://example.com/x", "http"),
			("http://127.0.0.1/", "non-public"),
			("http://10.0.0.5/", "non-public"),
			("http://169.254.169.254/latest/meta-data/", "non-public"),
		):
			with self.subTest(url), self.assertRaisesRegex(frappe.ValidationError, error):
				_validate_public_url(url)
		with patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 0))]):
			_validate_public_url("https://example.com/page")


class TestIngest(IntegrationTestCase):
	def setUp(self):
		store.drop_table()
		self.model = _make_model()
		_set_settings(
			embedding_model=self.model.name,
			embedding_dimension=DIM,
			chunk_size=60,
			chunk_overlap=10,
		)
		self.kb = frappe.get_doc({"doctype": "Flow Knowledge Base", "title": "Ingest KB"}).insert()

	def tearDown(self):
		store.drop_table()
		frappe.db.rollback()

	def _make_source(self, **values):
		values.setdefault("source_type", "Text")
		values.setdefault("title", "Ingest Source")
		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion"):
			return frappe.get_doc(
				{"doctype": "Flow Knowledge Source", "knowledge_base": self.kb.name, **values}
			).insert()

	def _ingest(self, source_name, embedding=None):
		def fake(input, **kwargs):
			return _embedding_response([[0.1, 0.2, 0.3, 0.4]] * len(input))

		with (
			patch("litellm.embedding", side_effect=embedding or fake),
			patch.object(frappe.db, "commit"),
			patch.object(frappe.db, "rollback"),
		):
			ingest_source(source_name)

	def _chunks(self, source_name):
		return frappe.get_all(
			"Flow Knowledge Chunk",
			filters={"source": source_name},
			fields=["name", "chunk_index", "knowledge_base"],
			order_by="chunk_index asc",
		)

	def _lance_count(self):
		return store._open_table().count_rows() if store.table_exists() else 0

	def test_a_knowledge_base_is_renamed_with_its_chunks_once_no_source_processes(self):
		store.ensure_table_for_dimension(DIM)
		store.add([{**row, "kb": self.kb.name} for row in _rows()])
		source = self._make_source(content="Wifi keeps dropping every afternoon.")
		source.db_set("status", "Processing")
		with self.assertRaisesRegex(frappe.ValidationError, "still processing"):
			frappe.rename_doc("Flow Knowledge Base", self.kb.name, "Ingest KB Renamed")

		source.db_set("status", "Completed")
		frappe.rename_doc("Flow Knowledge Base", self.kb.name, "Ingest KB Renamed", rebuild_search=False)
		hits = store.search([1.0, 0.0, 0.0, 0.0], kbs=["Ingest KB Renamed"], limit=5)
		self.assertEqual(sorted(hit["id"] for hit in hits), [1, 2, 3])

	def test_a_source_is_indexed_reindexed_and_purged_in_both_stores(self):
		source = self._make_source(content="Wifi keeps dropping on the office network every afternoon. " * 4)
		self._ingest(source.name)

		source.reload()
		self.assertEqual(source.status, "Completed")
		chunks = self._chunks(source.name)
		self.assertGreater(len(chunks), 1)
		self.assertEqual(source.chunk_count, len(chunks))
		self.assertEqual([c["chunk_index"] for c in chunks], list(range(len(chunks))))
		self.assertTrue(all(c["knowledge_base"] == self.kb.name for c in chunks))
		table = store._open_table()
		self.assertEqual(table.count_rows(), len(chunks))
		self.assertEqual(set(table.to_arrow().column("id").to_pylist()), {int(c["name"]) for c in chunks})

		# a resync replaces the chunks instead of adding to them
		self._ingest(source.name)
		again = self._chunks(source.name)
		self.assertEqual(len(again), len(chunks))
		self.assertEqual(self._lance_count(), len(again))
		self.assertEqual(len({c["name"] for c in again}), len(again))

		# a resync that cannot embed keeps the chunks it had, still searchable
		with self.assertRaises(frappe.ValidationError):
			self._ingest(source.name, embedding=RuntimeError("provider down"))
		self.assertEqual(self._chunks(source.name), again)
		self.assertEqual(self._lance_count(), len(again))

		source.delete()
		self.assertEqual(frappe.db.count("Flow Knowledge Chunk", {"source": source.name}), 0)
		self.assertEqual(self._lance_count(), 0)

	def test_an_empty_source_completes_and_a_failing_one_is_marked_failed(self):
		empty = self._make_source(content="   ")
		self._ingest(empty.name)
		empty.reload()
		self.assertEqual(empty.status, "Completed")
		self.assertEqual(empty.chunk_count, 0)
		self.assertEqual(self._chunks(empty.name), [])

		failing = self._make_source(source_type="URL", url="http://127.0.0.1/secret")
		with self.assertRaises(frappe.ValidationError):
			self._ingest(failing.name)
		failing.reload()
		self.assertEqual(failing.status, "Failed")
		self.assertIn("non-public", failing.error_log or "")

	def test_a_source_is_queued_for_indexing_on_insert_and_on_resync(self):
		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion") as enqueue:
			source = frappe.get_doc(
				{
					"doctype": "Flow Knowledge Source",
					"knowledge_base": self.kb.name,
					"source_type": "Text",
					"title": "Auto Ingest",
					"content": "indexed on creation",
				}
			).insert()
			enqueue.assert_called_once_with(source.name)
			enqueue.reset_mock()
			source.resync(rebuild=True)
		enqueue.assert_called_once_with(source.name, rebuild=True)

	def test_a_source_can_override_the_chunk_size(self):
		content = "Wifi keeps dropping on the office network every afternoon. " * 4
		default_source = self._make_source(content=content, title="Default Chunking")
		self._ingest(default_source.name)
		big_source = self._make_source(content=content, title="Big Chunking", chunk_size=400)
		self._ingest(big_source.name)
		self.assertGreater(len(self._chunks(default_source.name)), len(self._chunks(big_source.name)))

		with self.assertRaisesRegex(frappe.ValidationError, "Chunk Overlap"):
			self._make_source(content="x", chunk_size=100, chunk_overlap=100)

	def test_knowledge_needs_an_embedding_model_first(self):
		_set_settings(embedding_model="")
		with (
			self.subTest("knowledge base"),
			self.assertRaisesRegex(frappe.ValidationError, "Flow Knowledge Settings"),
		):
			frappe.get_doc({"doctype": "Flow Knowledge Base", "title": "No Model KB"}).insert()
		with (
			self.subTest("source"),
			self.assertRaisesRegex(frappe.ValidationError, "Flow Knowledge Settings"),
		):
			self._make_source(content="needs a model")


class TestDoctypeSync(IntegrationTestCase):
	def setUp(self):
		store.drop_table()
		self.model = _make_model()
		_set_settings(
			embedding_model=self.model.name,
			embedding_dimension=DIM,
			chunk_size=200,
			chunk_overlap=20,
		)
		self.kb = frappe.get_doc({"doctype": "Flow Knowledge Base", "title": "Sync KB"}).insert()

	def tearDown(self):
		store.drop_table()
		frappe.db.rollback()

	def _todo(self, description, **values):
		return frappe.get_doc({"doctype": "ToDo", "description": description, **values}).insert()

	def _source(self, filters=None):
		filters = filters or {"description": ["like", "SYNCTEST%"]}
		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion"):
			return frappe.get_doc(
				{
					"doctype": "Flow Knowledge Source",
					"knowledge_base": self.kb.name,
					"source_type": "DocType",
					"title": "ToDo Source",
					"reference_doctype": "ToDo",
					"content_fields": "description",
					"filters": json.dumps(filters),
				}
			).insert()

	def _sync(self, source_name, rebuild=False):
		def fake(input, **kwargs):
			return _embedding_response([[0.1, 0.2, 0.3, 0.4]] * len(input))

		embed = patch("litellm.embedding", side_effect=fake)
		with embed as mock, patch.object(frappe.db, "commit"), patch.object(frappe.db, "rollback"):
			ingest_source(source_name, rebuild=rebuild)
		return mock

	def _chunks(self, source_name, reference_name=None):
		filters = {"source": source_name}
		if reference_name:
			filters["reference_name"] = reference_name
		return frappe.get_all(
			"Flow Knowledge Chunk",
			filters=filters,
			fields=["name", "reference_name", "content_hash"],
			order_by="name asc",
		)

	def _references(self, source_name):
		return {c["reference_name"] for c in self._chunks(source_name)}

	def _lance_ids(self):
		if not store.table_exists():
			return set()
		return set(store._open_table().to_arrow().column("id").to_pylist())

	def _assert_lance_mirrors_chunks(self, source_name):
		self.assertEqual(self._lance_ids(), {int(c["name"]) for c in self._chunks(source_name)})

	def _content(self, source_name, reference_name):
		return frappe.get_all(
			"Flow Knowledge Chunk",
			filters={"source": source_name, "reference_name": reference_name},
			pluck="content",
		)

	def test_a_sweep_indexes_new_and_changed_rows_only(self):
		printer = self._todo("SYNCTEST printer is offline again", priority="Low")
		wifi = self._todo("SYNCTEST wifi keeps dropping on the office network")
		self._todo("UNRELATED should not be indexed")
		src = self._source()
		self._sync(src.name)

		src.reload()
		self.assertEqual(src.status, "Completed")
		self.assertIsNotNone(src.last_synced_at)
		chunks = self._chunks(src.name)
		self.assertEqual({c["reference_name"] for c in chunks}, {printer.name, wifi.name})
		self.assertTrue(all(c["content_hash"] for c in chunks))
		self._assert_lance_mirrors_chunks(src.name)

		# a new row is added, the indexed ones are kept
		scanner = self._todo("SYNCTEST second ticket about scanners")
		self._sync(src.name)
		self.assertEqual(self._references(src.name), {printer.name, wifi.name, scanner.name})
		self.assertTrue({c["name"] for c in chunks}.issubset({c["name"] for c in self._chunks(src.name)}))
		self._assert_lance_mirrors_chunks(src.name)

		# a content edit replaces the row's chunks
		before = self._chunks(src.name, wifi.name)
		frappe.db.set_value(
			"ToDo",
			wifi.name,
			"description",
			"SYNCTEST wifi fixed by replacing the router",
			update_modified=True,
		)
		self._sync(src.name)
		after = self._chunks(src.name, wifi.name)
		self.assertNotEqual(after[0]["content_hash"], before[0]["content_hash"])
		self.assertTrue({c["name"] for c in before}.isdisjoint({c["name"] for c in after}))
		self._assert_lance_mirrors_chunks(src.name)

		# any other edit is skipped without embedding
		before = {c["name"]: c["content_hash"] for c in self._chunks(src.name, printer.name)}
		frappe.db.set_value("ToDo", printer.name, "priority", "High", update_modified=True)
		self.assertEqual(self._sync(src.name).call_count, 0)
		self.assertEqual({c["name"]: c["content_hash"] for c in self._chunks(src.name, printer.name)}, before)

	def test_rows_that_leave_the_source_lose_their_chunks(self):
		keep = self._todo("SYNCTEST keep this one")
		closed = self._todo("SYNCTEST printer offline")
		deleted = self._todo("SYNCTEST delete this one")
		raw = self._todo("SYNCTEST delete this one without a trace")
		src = self._source(filters={"description": ["like", "SYNCTEST%"], "status": "Open"})
		self._sync(src.name)
		self.assertEqual(self._references(src.name), {keep.name, closed.name, deleted.name, raw.name})

		# leaving the filter
		frappe.db.set_value("ToDo", closed.name, "status", "Cancelled", update_modified=True)
		self.assertEqual(self._sync(src.name).call_count, 0)
		self.assertEqual(self._chunks(src.name, closed.name), [])
		self._assert_lance_mirrors_chunks(src.name)

		# a delete leaves a tombstone the sweep reads
		frappe.delete_doc("ToDo", deleted.name)
		self._sync(src.name)
		self.assertEqual(self._references(src.name), {keep.name, raw.name})
		self._assert_lance_mirrors_chunks(src.name)

		# a raw delete leaves none: the sweep misses it, a reconcile purges it
		frappe.db.delete("ToDo", {"name": raw.name})
		self._sync(src.name)
		self.assertEqual(self._references(src.name), {keep.name, raw.name})
		reconcile_source(src.name)
		self.assertEqual(self._references(src.name), {keep.name})
		self._assert_lance_mirrors_chunks(src.name)
		src.reload()
		self.assertEqual(src.chunk_count, len(self._chunks(src.name)))

	def test_only_redefining_a_source_reindexes_it(self):
		todo = self._todo("SYNCTEST alpha", status="Closed")
		source = self._source()
		self._sync(source.name)
		self.assertNotIn("Status", " ".join(self._content(source.name, todo.name)))

		source.title = "Renamed source"
		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion") as enqueue:
			source.save()
		enqueue.assert_not_called()

		source.content_fields = "description, status"
		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion") as enqueue:
			source.save()
		enqueue.assert_called_once_with(source.name, rebuild=True)

		self._sync(source.name, rebuild=True)
		self.assertIn("Status: Closed", " ".join(self._content(source.name, todo.name)))

	def test_sync_due_sources_enqueues_only_auto_sync_doctypes(self):
		auto = self._source()
		auto.db_set("auto_sync", 1, update_modified=False)
		manual = self._source()
		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion"):
			text = frappe.get_doc(
				{
					"doctype": "Flow Knowledge Source",
					"knowledge_base": self.kb.name,
					"source_type": "Text",
					"title": "Text Source",
					"content": "not a doctype",
					"auto_sync": 1,
				}
			).insert()

		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion") as mock:
			sync_due_sources()
		enqueued = {call.args[0] for call in mock.call_args_list}

		self.assertIn(auto.name, enqueued)
		self.assertNotIn(manual.name, enqueued)
		self.assertNotIn(text.name, enqueued)


class TestRetriever(IntegrationTestCase):
	def setUp(self):
		store.drop_table()
		self.model = _make_model()
		_set_settings(
			embedding_model=self.model.name,
			embedding_dimension=DIM,
			chunk_size=200,
			chunk_overlap=20,
		)
		self.kb = frappe.get_doc({"doctype": "Flow Knowledge Base", "title": "Retrieve KB"}).insert()

	def tearDown(self):
		store.drop_table()
		frappe.db.rollback()

	def _fake_embed(self, input, **kwargs):
		return _embedding_response([[0.1, 0.2, 0.3, 0.4]] * len(input))

	def _ingest(self, source_name):
		with (
			patch("litellm.embedding", side_effect=self._fake_embed),
			patch.object(frappe.db, "commit"),
			patch.object(frappe.db, "rollback"),
		):
			ingest_source(source_name)

	def _text_source(self, kb, content, title="Doc"):
		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion"):
			source = frappe.get_doc(
				{
					"doctype": "Flow Knowledge Source",
					"knowledge_base": kb,
					"source_type": "Text",
					"title": title,
					"content": content,
				}
			).insert()
		self._ingest(source.name)
		return source

	def _retrieve(self, query, **kwargs):
		with patch("litellm.embedding", side_effect=self._fake_embed):
			return retrieve(query, **kwargs)

	def test_retrieval_searches_only_the_given_enabled_knowledge_bases(self):
		other = frappe.get_doc({"doctype": "Flow Knowledge Base", "title": "Other KB"}).insert()
		mine = self._text_source(self.kb.name, "laptop onboarding steps. " * 4)
		self._text_source(other.name, "laptop onboarding steps. " * 4, title="Other Doc")

		for kbs in ([self.kb.name], [self.kb.name, "No Such KB"]):
			with self.subTest(kbs=kbs):
				results = self._retrieve("laptop", kbs=kbs)
				self.assertTrue(results)
				self.assertTrue(all("laptop" in r["content"] for r in results))
				self.assertTrue(all(r["source"] == mine.title for r in results))

		self.assertEqual(self._retrieve("   ", kbs=[self.kb.name]), [])
		# an empty scope is refused, never read as every knowledge base
		with self.assertRaises(frappe.ValidationError):
			retrieve("laptop", kbs=[])
		frappe.db.set_value("Flow Knowledge Base", self.kb.name, "enabled", 0)
		self.assertEqual(self._retrieve("laptop", kbs=[self.kb.name]), [])

	def test_the_search_type_decides_whether_text_is_searched(self):
		for search_type, text in (("Vector", None), ("Hybrid", "laptop")):
			with self.subTest(search_type):
				frappe.db.set_single_value("Flow Knowledge Settings", "search_type", search_type)
				with patch("frappe.flow.knowledge.store.search", return_value=[]) as mock:
					self._retrieve("laptop", kbs=[self.kb.name])
				self.assertEqual(mock.call_args.kwargs["text"], text)


class TestKnowledgeBuilder(IntegrationTestCase):
	def setUp(self):
		store.drop_table()
		self.model = _make_model()
		_set_settings(
			embedding_model=self.model.name,
			embedding_dimension=DIM,
			chunk_size=200,
			chunk_overlap=20,
		)

	def tearDown(self):
		store.drop_table()
		frappe.db.rollback()

	def _fake_embed(self, input, **kwargs):
		return _embedding_response([[0.1, 0.2, 0.3, 0.4]] * len(input))

	def _add(self, add, *args):
		with (
			patch("litellm.embedding", side_effect=self._fake_embed),
			patch.object(frappe.db, "commit"),
			patch.object(frappe.db, "rollback"),
		):
			return add(*args)

	def _assert_searchable(self, kb, source_name):
		source = frappe.get_doc("Flow Knowledge Source", source_name)
		self.assertEqual(source.source_type, "Text")
		with patch("litellm.embedding", side_effect=self._fake_embed):
			results = retrieve("laptop", kbs=[kb.name])
		self.assertTrue(results)
		self.assertTrue(all(r["source"] == source.title for r in results))
		return source

	def test_added_text_and_local_files_are_searchable_on_return(self):
		import os
		import tempfile

		kb = Knowledge("Builder KB")
		self.assertEqual(Knowledge("Builder KB").name, kb.name)
		self.assertEqual(frappe.db.count("Flow Knowledge Base", {"title": "Builder KB"}), 1)
		with patch("frappe.flow.knowledge.ingest.enqueue_ingestion") as enqueue:
			source_name = self._add(kb.add_text, "laptop setup guide for new employees. " * 4)
		enqueue.assert_not_called()
		self.assertEqual(self._assert_searchable(kb, source_name).knowledge_base, kb.name)

		fd, path = tempfile.mkstemp(suffix=".txt")
		try:
			with os.fdopen(fd, "w") as handle:
				handle.write("laptop provisioning runbook for it staff. " * 4)
			local_kb = Knowledge("Local File KB")
			source_name = self._add(local_kb.add_local_file, path)
		finally:
			os.unlink(path)
		self.assertEqual(self._assert_searchable(local_kb, source_name).title, os.path.basename(path))

	def test_add_doctype_translates_fields_and_filters(self):
		kb = Knowledge("DocType KB")
		with patch("frappe.flow.knowledge.ingest.ingest_source") as ingest:
			source_name = kb.add_doctype(
				"ToDo",
				content_fields=["description", "status"],
				filters={"status": "Open"},
			)

		source = frappe.get_doc("Flow Knowledge Source", source_name)
		self.assertEqual(source.source_type, "DocType")
		self.assertEqual(source.reference_doctype, "ToDo")
		self.assertEqual(source.content_fields, "description, status")
		self.assertEqual(json.loads(source.filters), {"status": "Open"})
		ingest.assert_called_once_with(source_name)


class TestKnowledgeSearchTool(IntegrationTestCase):
	def setUp(self):
		_set_settings(embedding_model=_make_model().name)

	def tearDown(self):
		frappe.db.rollback()

	def _search_tool(self, knowledge):
		from frappe.flow.tools.builtins import bind_search_knowledge

		return bind_search_knowledge([kb.name for kb in knowledge])

	def test_the_tool_takes_only_a_query_and_searches_its_bound_bases(self):
		first = Knowledge("KB One")
		second = Knowledge("KB Two")
		search = self._search_tool([first, second])
		self.assertEqual(set(search.parameters["properties"]), {"query"})
		with patch("frappe.flow.knowledge.retriever.retrieve", return_value=[]) as mock:
			search(query="hi", ctx=CTX)
		mock.assert_called_once_with("hi", kbs=[first.name, second.name])

	def test_the_description_lists_bound_bases_that_have_one(self):
		kb = Knowledge("Travel KB", description="Company travel and reimbursement policy.")
		search = self._search_tool([kb])
		self.assertIn("Travel KB", search.description)
		self.assertIn("Company travel and reimbursement policy.", search.description)

		search = self._search_tool([Knowledge("Plain KB")])
		self.assertNotIn("This agent's knowledge bases:", search.description)


class TestAttachmentStore(IntegrationTestCase):
	def setUp(self):
		attachment_store.drop_table()

	def tearDown(self):
		attachment_store.drop_table()

	def _seed(self):
		attachment_store.ensure_table(DIM)
		attachment_store.add(
			"SES-a",
			"ATT-1",
			["wifi not connecting on macbook", "printer not detected on windows"],
			[[1.0, 0.0, 0.0, 0.0], [0.0, 1.0, 0.0, 0.0]],
		)
		attachment_store.add("SES-b", "ATT-2", ["how to set up your laptop"], [[0.9, 0.1, 0.0, 0.0]])

	def test_the_table_is_recreated_for_a_new_dimension(self):
		self.assertEqual(attachment_store.search([1.0, 0.0, 0.0, 0.0], session="SES-a"), [])
		with self.assertRaises(ValueError):
			attachment_store.ensure_table(0)

		attachment_store.ensure_table(DIM)
		attachment_store.ensure_table(DIM)
		self.assertTrue(attachment_store._table_exists())
		with self.assertRaises(ValueError):
			attachment_store.add("SES-a", "ATT-1", ["one", "two"], [[1.0, 0.0, 0.0, 0.0]])

		self._seed()
		attachment_store.ensure_table(DIM + 1)  # different width -> dropped and recreated
		self.assertEqual(attachment_store._open_table().count_rows(), 0)

	def test_search_is_scoped_to_one_session(self):
		self._seed()
		with self.assertRaises(ValueError):
			attachment_store.search([1.0, 0.0, 0.0, 0.0], session="")

		hits = attachment_store.search([1.0, 0.0, 0.0, 0.0], session="SES-a", limit=5)
		self.assertTrue(hits)
		self.assertTrue(all("laptop" not in h["content"] for h in hits))  # SES-b not leaked
		hits = attachment_store.search([0.9, 0.1, 0.0, 0.0], session="SES-b", limit=5)
		self.assertEqual(
			[(h["attachment"], h["content"]) for h in hits], [("ATT-2", "how to set up your laptop")]
		)

		hits = attachment_store.search([1.0, 0.0, 0.0, 0.0], session="SES-a", text="printer", limit=2)
		self.assertTrue(any("printer" in h["content"] for h in hits))

	def test_delete_by_session_or_attachment(self):
		with self.assertRaises(ValueError):
			attachment_store.delete()
		for criteria, emptied, kept in (
			({"session": "SES-a", "attachment": "ATT-1"}, ["SES-a"], []),
			({"session": "SES-b"}, ["SES-b"], []),
			({"session": ["SES-a", "SES-b"]}, ["SES-a", "SES-b"], []),
			({"session": ["SES-a"]}, ["SES-a"], ["SES-b"]),
			({"session": []}, [], ["SES-b"]),
		):
			with self.subTest(criteria=criteria):
				attachment_store.drop_table()
				self._seed()
				attachment_store.delete(**criteria)
				for session in emptied:
					self.assertEqual(
						attachment_store.search([1.0, 0.0, 0.0, 0.0], session=session, limit=5), []
					)
				for session in kept:
					self.assertTrue(attachment_store.search([1.0, 0.0, 0.0, 0.0], session=session, limit=5))


class TestRetrieveAttachments(IntegrationTestCase):
	def setUp(self):
		attachment_store.drop_table()
		self.model = _make_model()
		_set_settings(embedding_model=self.model.name, embedding_dimension=DIM, search_type="Hybrid")
		attachment_store.ensure_table(DIM)
		attachment_store.add(
			"SES-x",
			"ATT-1",
			["the refund policy allows returns within thirty days"],
			[[1.0, 0.0, 0.0, 0.0]],
		)

	def tearDown(self):
		attachment_store.drop_table()
		frappe.db.rollback()

	def _fake_embed(self, input, **kwargs):
		return _embedding_response([[1.0, 0.0, 0.0, 0.0]] * len(input))

	def test_retrieval_is_scoped_to_the_session(self):
		with patch("litellm.embedding", side_effect=self._fake_embed):
			results = retrieve_attachments("refund", session="SES-x")
			self.assertEqual(retrieve_attachments("refund", session="SES-other"), [])
		self.assertTrue(results)
		self.assertIn("refund", results[0]["content"])

		self.assertEqual(retrieve_attachments("refund", session=""), [])
		with patch("litellm.embedding") as mocked:
			self.assertEqual(retrieve_attachments("   ", session="SES-x"), [])
		mocked.assert_not_called()  # no embedding call for an empty query

	def test_the_search_type_decides_whether_text_is_searched(self):
		for search_type, text in (("Vector", None), ("Hybrid", "refund")):
			with self.subTest(search_type):
				_set_settings(search_type=search_type)
				with (
					patch("litellm.embedding", side_effect=self._fake_embed),
					patch.object(attachment_store, "search", return_value=[]) as mocked,
				):
					retrieve_attachments("refund", session="SES-x")
				self.assertEqual(mocked.call_args.kwargs["text"], text)


class TestSystemGeneratedProtection(IntegrationTestCase):
	"""Desk edits can't break app-owned system rows; the app's sync (ignore_permissions) can."""

	def setUp(self):
		self.model = _make_model()
		_set_settings(embedding_model=self.model.name, embedding_dimension=DIM)

	def tearDown(self):
		frappe.db.rollback()

	def _kb(self, system=True, title="Builtin KB"):
		return frappe.get_doc(
			{"doctype": "Flow Knowledge Base", "title": title, "is_system_generated": int(system)}
		).insert()

	def _source(self, kb, system=True, **values):
		values.setdefault("source_type", "Text")
		values.setdefault("title", "Builtin Source")
		values.setdefault("content", "some indexed content")
		doc = frappe.get_doc(
			{
				"doctype": "Flow Knowledge Source",
				"knowledge_base": kb.name,
				"is_system_generated": int(system),
				**values,
			}
		)
		doc.flags.skip_auto_ingest = True
		return doc.insert()

	def test_a_system_generated_kb_is_deleted_only_by_its_app_and_never_renamed(self):
		from frappe.model.rename_doc import rename_doc

		kb = self._kb()
		with (
			self.subTest("delete"),
			self.assertRaisesRegex(frappe.ValidationError, "Cannot delete system-generated"),
		):
			frappe.delete_doc("Flow Knowledge Base", kb.name)
		with (
			self.subTest("rename"),
			self.assertRaisesRegex(frappe.ValidationError, "Cannot rename system-generated"),
		):
			frappe.rename_doc("Flow Knowledge Base", kb.name, "Renamed KB")
		with (
			self.subTest("rename as the app"),
			self.assertRaisesRegex(frappe.ValidationError, "Cannot rename system-generated"),
		):
			rename_doc("Flow Knowledge Base", kb.name, "Renamed KB", ignore_permissions=True)
		with (
			self.subTest("unset the flag"),
			self.assertRaisesRegex(frappe.ValidationError, "Cannot remove the system-generated flag"),
		):
			frappe.get_doc("Flow Knowledge Base", kb.name).update({"is_system_generated": 0}).save()

		frappe.delete_doc("Flow Knowledge Base", kb.name, ignore_permissions=True)
		self.assertFalse(frappe.db.exists("Flow Knowledge Base", kb.name))
		regular = self._kb(system=False, title="User KB")
		frappe.delete_doc("Flow Knowledge Base", regular.name)
		self.assertFalse(frappe.db.exists("Flow Knowledge Base", regular.name))

	def test_desk_may_edit_only_the_content_of_a_system_generated_source(self):
		source = self._source(self._kb())

		def saved():
			return frappe.get_doc("Flow Knowledge Source", source.name)

		with (
			self.subTest("delete"),
			self.assertRaisesRegex(frappe.ValidationError, "Cannot delete system-generated"),
		):
			frappe.delete_doc("Flow Knowledge Source", source.name)
		with (
			self.subTest("change type"),
			self.assertRaisesRegex(frappe.ValidationError, "Cannot change Source Type"),
		):
			saved().update({"source_type": "URL", "url": "https://example.com"}).save()
		with (
			self.subTest("unset the flag"),
			self.assertRaisesRegex(frappe.ValidationError, "Cannot remove the system-generated flag"),
		):
			saved().update({"is_system_generated": 0}).save()

		saved().update({"content": "refreshed content"}).save()
		self.assertEqual(
			frappe.db.get_value("Flow Knowledge Source", source.name, "content"), "refreshed content"
		)
		saved().update({"source_type": "URL", "url": "https://example.com"}).save(ignore_permissions=True)
		self.assertEqual(frappe.db.get_value("Flow Knowledge Source", source.name, "source_type"), "URL")
